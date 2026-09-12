from pathlib import Path
import re

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


# ---------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------

PARTICIPANT_RESULTS_DIR = Path("participant_results")
ANSWER_KEY_FILE = Path("answer_key.csv")
OUTPUT_DIR = Path("study_analysis")

QUESTION_COLUMN = "Question"
PARTICIPANT_ANSWER_COLUMN = "Your Answer"
ANSWER_KEY_COLUMN = "Correct Answer"

HEADER_BLUE = "1F4E78"
LIGHT_GREEN = "C6EFCE"


# ---------------------------------------------------------------------
# Utility functions
# ---------------------------------------------------------------------

def question_number(question_id):
    """Convert Q12 or 12 into integer 12 for sorting."""
    match = re.fullmatch(r"\s*Q?(\d+)\s*", str(question_id))

    if match is None:
        raise ValueError(
            f"Invalid question label '{question_id}'. "
            "Expected a label such as Q1, Q2, ..., Q20."
        )

    return int(match.group(1))


def normalise_question(value):
    """Convert question values into standard form: Q1, Q2, ..."""
    return f"Q{question_number(value)}"


def normalise_option(value):
    """Convert A, B, Option A, or Option B into A or B."""
    text = str(value).strip().upper()

    if text in {"A", "OPTION A"}:
        return "A"

    if text in {"B", "OPTION B"}:
        return "B"

    raise ValueError(
        f"Invalid option value '{value}'. "
        "Expected A, B, Option A, or Option B."
    )


def require_columns(dataframe, required_columns, file_label):
    """Check that a CSV includes all required columns."""
    missing_columns = [
        column
        for column in required_columns
        if column not in dataframe.columns
    ]

    if missing_columns:
        raise ValueError(
            f"{file_label} is missing required column(s): "
            f"{missing_columns}\n"
            f"Available columns: {list(dataframe.columns)}"
        )


def save_dataframe(dataframe, filename):
    """Save a dataframe as a CSV without its index."""
    dataframe.to_csv(
        OUTPUT_DIR / filename,
        index=False,
    )


def add_value_labels(axis, bars, value_format="{:.1f}%"):
    """Write a value label above each bar."""
    for bar in bars:
        height = bar.get_height()

        axis.annotate(
            value_format.format(height),
            xy=(
                bar.get_x() + bar.get_width() / 2,
                height,
            ),
            xytext=(0, 4),
            textcoords="offset points",
            ha="center",
            va="bottom",
            fontsize=10,
        )


# ---------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------

def style_excel_sheet(worksheet, table_name):
    """Apply readable formatting and add an Excel table."""
    header_fill = PatternFill(
        fill_type="solid",
        fgColor=HEADER_BLUE,
    )

    for cell in worksheet[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row,
    ):
        for cell in row:
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 32

    for column_cells in worksheet.columns:
        column_letter = column_cells[0].column_letter

        maximum_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )

        worksheet.column_dimensions[column_letter].width = min(
            max(maximum_length + 3, 12),
            28,
        )

    if worksheet.max_row >= 2:
        table = Table(
            displayName=table_name,
            ref=worksheet.dimensions,
        )

        table_style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        table.tableStyleInfo = table_style
        worksheet.add_table(table)


def export_response_matrix_excel(
    response_matrix_export,
    participant_columns,
):
    """
    Export the response matrix.

    Participant response cells are highlighted light green when their
    selected option matches the answer key for that question.
    """
    output_path = OUTPUT_DIR / "response_matrix.xlsx"

    response_matrix_export.to_excel(
        output_path,
        sheet_name="Response Matrix",
        index=False,
    )

    workbook = load_workbook(output_path)
    worksheet = workbook["Response Matrix"]

    style_excel_sheet(
        worksheet,
        "ResponseMatrixTable",
    )

    correct_fill = PatternFill(
        fill_type="solid",
        fgColor=LIGHT_GREEN,
    )

    header_lookup = {
        cell.value: cell.column
        for cell in worksheet[1]
    }

    correct_option_column = header_lookup["Correct Option"]

    for row_index in range(2, worksheet.max_row + 1):
        correct_option = worksheet.cell(
            row=row_index,
            column=correct_option_column,
        ).value

        for participant in participant_columns:
            participant_column = header_lookup[participant]

            selected_option = worksheet.cell(
                row=row_index,
                column=participant_column,
            ).value

            if selected_option == correct_option:
                worksheet.cell(
                    row=row_index,
                    column=participant_column,
                ).fill = correct_fill

    workbook.save(output_path)


# ---------------------------------------------------------------------
# Input loading
# ---------------------------------------------------------------------

def load_answer_key():
    """
    Load answer_key.csv.

    Required columns:
    - Question
    - Correct Answer
    """
    if not ANSWER_KEY_FILE.exists():
        raise FileNotFoundError(
            f"Answer key not found: {ANSWER_KEY_FILE}"
        )

    answer_key = pd.read_csv(ANSWER_KEY_FILE)

    require_columns(
        answer_key,
        [QUESTION_COLUMN, ANSWER_KEY_COLUMN],
        f"Answer key '{ANSWER_KEY_FILE}'",
    )

    answer_key = answer_key.copy()

    answer_key[QUESTION_COLUMN] = answer_key[
        QUESTION_COLUMN
    ].apply(
        normalise_question
    )

    answer_key["correct_option"] = answer_key[
        ANSWER_KEY_COLUMN
    ].apply(
        normalise_option
    )

    if answer_key[QUESTION_COLUMN].duplicated().any():
        duplicate_questions = answer_key.loc[
            answer_key[QUESTION_COLUMN].duplicated(),
            QUESTION_COLUMN,
        ].tolist()

        raise ValueError(
            f"Answer key contains duplicate questions: "
            f"{duplicate_questions}"
        )

    return answer_key[
        [
            QUESTION_COLUMN,
            "correct_option",
        ]
    ]


def load_participant_results():
    """
    Load any number of Study 2 participant CSV files.

    Required columns in every participant CSV:
    - Question
    - Your Answer
    """
    if not PARTICIPANT_RESULTS_DIR.exists():
        raise FileNotFoundError(
            f"Participant results folder not found: "
            f"{PARTICIPANT_RESULTS_DIR}"
        )

    participant_files = sorted(
        PARTICIPANT_RESULTS_DIR.glob("*.csv")
    )

    if not participant_files:
        raise FileNotFoundError(
            f"No participant CSV files found in: "
            f"{PARTICIPANT_RESULTS_DIR}"
        )

    all_participant_frames = []

    for file_path in participant_files:
        participant_id = file_path.stem
        participant_df = pd.read_csv(file_path)

        require_columns(
            participant_df,
            [QUESTION_COLUMN, PARTICIPANT_ANSWER_COLUMN],
            f"Participant file '{file_path.name}'",
        )

        participant_df = participant_df.copy()

        participant_df[QUESTION_COLUMN] = participant_df[
            QUESTION_COLUMN
        ].apply(
            normalise_question
        )

        participant_df["selected_option"] = participant_df[
            PARTICIPANT_ANSWER_COLUMN
        ].apply(
            normalise_option
        )

        if participant_df[QUESTION_COLUMN].duplicated().any():
            duplicate_questions = participant_df.loc[
                participant_df[QUESTION_COLUMN].duplicated(),
                QUESTION_COLUMN,
            ].tolist()

            raise ValueError(
                f"Participant file '{file_path.name}' has duplicate "
                f"questions: {duplicate_questions}"
            )

        participant_df["participant"] = participant_id

        all_participant_frames.append(
            participant_df[
                [
                    QUESTION_COLUMN,
                    "participant",
                    "selected_option",
                ]
            ]
        )

    return pd.concat(
        all_participant_frames,
        ignore_index=True,
    )


# ---------------------------------------------------------------------
# Plot functions
# ---------------------------------------------------------------------

def create_participant_accuracy_plot(participant_summary):
    """Create participant agreement plot with y-axis extending to 110%."""
    figure, axis = plt.subplots(
        figsize=(max(8, len(participant_summary) * 1.3), 5.2),
        layout="constrained",
    )

    bars = axis.bar(
        participant_summary["Participant"],
        participant_summary["Accuracy (%)"],
        color="tab:blue",
    )

    axis.set_ylim(0, 110)
    axis.set_ylabel("Agreement with Answer Key (%)")
    axis.set_xlabel("Participant")
    axis.set_title(
        "Participant Agreement with Answer Key",
        fontweight="bold",
    )

    axis.grid(
        axis="y",
        linestyle="--",
        alpha=0.4,
    )

    add_value_labels(axis, bars)

    figure.savefig(
        OUTPUT_DIR / "participant_accuracy.png",
        dpi=200,
        bbox_inches="tight",
        facecolor="white",
    )

    plt.close(figure)


def create_question_accuracy_plot(question_summary):
    """Create question-level agreement-with-key plot."""
    figure, axis = plt.subplots(
        figsize=(14, 5.2),
        layout="constrained",
    )

    colors = [
        "tab:green"
        if accuracy >= 75
        else "tab:orange"
        if accuracy >= 50
        else "tab:red"
        for accuracy in question_summary["Agreement with Answer Key (%)"]
    ]

    bars = axis.bar(
        question_summary["Question"],
        question_summary["Agreement with Answer Key (%)"],
        color=colors,
    )

    axis.set_ylim(0, 110)
    axis.set_ylabel("Agreement with Answer Key (%)")
    axis.set_xlabel("Question")
    axis.set_title(
        "Question-Level Agreement with Answer Key",
        fontweight="bold",
    )

    axis.grid(
        axis="y",
        linestyle="--",
        alpha=0.4,
    )

    add_value_labels(axis, bars)

    figure.savefig(
        OUTPUT_DIR / "question_accuracy.png",
        dpi=200,
        bbox_inches="tight",
        facecolor="white",
    )

    plt.close(figure)


def create_option_selection_plot(question_summary):
    """
    Create a grouped bar chart showing how many participants chose
    Option A and Option B for each question.

    The majority response, A, B, or Tie, is shown above each pair of bars.
    """
    questions = question_summary["Question"].tolist()
    selected_a = question_summary["Selected A"].to_numpy()
    selected_b = question_summary["Selected B"].to_numpy()
    majority_choice = question_summary["Majority Choice"].tolist()

    number_of_questions = len(questions)
    x_positions = np.arange(number_of_questions)
    bar_width = 0.36

    figure, axis = plt.subplots(
        figsize=(max(14, number_of_questions * 0.75), 6.4),
        layout="constrained",
    )

    bars_a = axis.bar(
        x_positions - bar_width / 2,
        selected_a,
        width=bar_width,
        color="tab:orange",
        label="Selected A",
    )

    bars_b = axis.bar(
        x_positions + bar_width / 2,
        selected_b,
        width=bar_width,
        color="tab:blue",
        label="Selected B",
    )

    axis.bar_label(
        bars_a,
        padding=3,
        fontsize=9,
    )

    axis.bar_label(
        bars_b,
        padding=3,
        fontsize=9,
    )

    maximum_selection_count = max(
        int(np.max(selected_a)),
        int(np.max(selected_b)),
    )

    axis.set_ylim(0, maximum_selection_count + 1.9)

    for x_position, majority, count_a, count_b in zip(
        x_positions,
        majority_choice,
        selected_a,
        selected_b,
    ):
        axis.text(
            x_position,
            max(count_a, count_b) + 0.55,
            f"Majority: {majority}",
            ha="center",
            va="bottom",
            fontsize=6.5,
            fontweight="bold",
            color="black",
        )

    axis.set_xticks(x_positions)
    axis.set_xticklabels(questions)
    axis.set_xlabel("Question")
    axis.set_ylabel("Number of Participants")
    axis.set_title(
        "Option Selections by Question",
        fontweight="bold",
    )

    axis.legend(
        loc="upper right",
        frameon=True,
    )

    axis.grid(
        axis="y",
        linestyle="--",
        alpha=0.4,
    )

    figure.savefig(
        OUTPUT_DIR / "option_selection_by_question.png",
        dpi=200,
        bbox_inches="tight",
        facecolor="white",
    )

    plt.close(figure)


def create_response_heatmap(
    response_matrix,
    correct_options_by_question,
):
    """
    Create response matrix heatmap.

    Light green means selected option matches the answer key.
    White means selected option does not match the answer key.
    The participant's A/B selection is shown inside each cell.
    """
    participants = list(response_matrix.columns)
    questions = list(response_matrix.index)

    correctness_matrix = np.zeros(
        response_matrix.shape,
        dtype=float,
    )

    for question_index, question in enumerate(questions):
        correct_option = correct_options_by_question.loc[question]

        for participant_index, participant in enumerate(participants):
            selected_option = response_matrix.loc[
                question,
                participant,
            ]

            correctness_matrix[
                question_index,
                participant_index,
            ] = int(selected_option == correct_option)

    figure_width = max(8, len(participants) * 1.25 + 4)
    figure_height = max(8, len(questions) * 0.42 + 2)

    figure, axis = plt.subplots(
        figsize=(figure_width, figure_height),
        layout="constrained",
    )

    cmap = plt.matplotlib.colors.ListedColormap(
        ["#FFFFFF", "#C6EFCE"]
    )

    axis.imshow(
        correctness_matrix,
        cmap=cmap,
        vmin=0,
        vmax=1,
        aspect="auto",
    )

    axis.set_xticks(range(len(participants)))
    axis.set_xticklabels(
        participants,
        rotation=35,
        ha="right",
    )

    axis.set_yticks(range(len(questions)))
    axis.set_yticklabels(questions)

    axis.set_xlabel("Participant")
    axis.set_ylabel("Question")
    axis.set_title(
        "Participant Response Matrix",
        fontweight="bold",
    )

    for question_index, question in enumerate(questions):
        for participant_index, participant in enumerate(participants):
            selected_option = response_matrix.loc[
                question,
                participant,
            ]

            axis.text(
                participant_index,
                question_index,
                selected_option,
                ha="center",
                va="center",
                fontsize=10,
                fontweight="bold",
                color="black",
            )

    legend_handles = [
        plt.matplotlib.patches.Patch(
            facecolor="#C6EFCE",
            edgecolor="#9BBB59",
            label="Matches answer key",
        ),
        plt.matplotlib.patches.Patch(
            facecolor="#FFFFFF",
            edgecolor="#BFBFBF",
            label="Does not match answer key",
        ),
    ]

    axis.legend(
        handles=legend_handles,
        loc="upper center",
        bbox_to_anchor=(0.5, -0.07),
        ncol=2,
        frameon=False,
    )

    figure.savefig(
        OUTPUT_DIR / "response_heatmap.png",
        dpi=200,
        bbox_inches="tight",
        facecolor="white",
    )

    plt.close(figure)


def create_pairwise_agreement_heatmap(pairwise_agreement):
    """Create participant-to-participant response agreement heatmap."""
    values = pairwise_agreement.to_numpy(dtype=float)
    participants = list(pairwise_agreement.columns)

    figure_width = max(7, len(participants) * 1.2 + 3)
    figure_height = max(6, len(participants) * 1.0 + 2)

    figure, axis = plt.subplots(
        figsize=(figure_width, figure_height),
        layout="constrained",
    )

    image = axis.imshow(
        values,
        cmap="Blues",
        vmin=0,
        vmax=100,
        aspect="auto",
    )

    axis.set_xticks(range(len(participants)))
    axis.set_xticklabels(
        participants,
        rotation=35,
        ha="right",
    )

    axis.set_yticks(range(len(participants)))
    axis.set_yticklabels(participants)

    axis.set_xlabel("Participant")
    axis.set_ylabel("Participant")
    axis.set_title(
        "Pairwise Participant Agreement",
        fontweight="bold",
    )

    for row_index in range(values.shape[0]):
        for column_index in range(values.shape[1]):
            value = values[row_index, column_index]

            if np.isnan(value):
                label = ""
                text_color = "black"
            else:
                label = f"{value:.1f}%"
                text_color = "white" if value >= 60 else "black"

            axis.text(
                column_index,
                row_index,
                label,
                ha="center",
                va="center",
                fontsize=10,
                color=text_color,
            )

    colorbar = figure.colorbar(
        image,
        ax=axis,
        shrink=0.85,
    )

    colorbar.ax.set_ylabel(
        "Agreement (%)",
        rotation=270,
        labelpad=16,
    )

    figure.savefig(
        OUTPUT_DIR / "pairwise_agreement_heatmap.png",
        dpi=200,
        bbox_inches="tight",
        facecolor="white",
    )

    plt.close(figure)


# ---------------------------------------------------------------------
# Main analysis
# ---------------------------------------------------------------------

def main():
    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    answer_key = load_answer_key()
    responses = load_participant_results()

    participant_ids = sorted(
        responses["participant"].unique()
    )

    expected_questions = set(
        answer_key[QUESTION_COLUMN]
    )

    for participant in participant_ids:
        answered_questions = set(
            responses.loc[
                responses["participant"] == participant,
                QUESTION_COLUMN,
            ]
        )

        missing_questions = sorted(
            expected_questions - answered_questions,
            key=question_number,
        )

        extra_questions = sorted(
            answered_questions - expected_questions,
            key=question_number,
        )

        if missing_questions or extra_questions:
            raise ValueError(
                f"Question mismatch for participant '{participant}'.\n"
                f"Missing: {missing_questions}\n"
                f"Unexpected: {extra_questions}"
            )

    merged = responses.merge(
        answer_key,
        on=QUESTION_COLUMN,
        how="left",
        validate="many_to_one",
    )

    if merged["correct_option"].isna().any():
        missing_key_questions = merged.loc[
            merged["correct_option"].isna(),
            QUESTION_COLUMN,
        ].unique()

        raise ValueError(
            "Questions missing from answer_key.csv: "
            f"{list(missing_key_questions)}"
        )

    merged["matches_answer_key"] = (
        merged["selected_option"] == merged["correct_option"]
    )

    merged["Question Number"] = merged[
        QUESTION_COLUMN
    ].apply(
        question_number
    )

    merged = merged.sort_values(
        ["participant", "Question Number"],
    ).reset_index(drop=True)

    # -----------------------------------------------------------------
    # 1. Merged participant answers
    # -----------------------------------------------------------------

    all_responses_export = merged[
        [
            "participant",
            QUESTION_COLUMN,
            "selected_option",
            "correct_option",
            "matches_answer_key",
        ]
    ].rename(
        columns={
            "participant": "Participant",
            QUESTION_COLUMN: "Question",
            "selected_option": "Selected Option",
            "correct_option": "Correct Option",
            "matches_answer_key": "Matches Answer Key",
        }
    )

    save_dataframe(
        all_responses_export,
        "all_responses_merged.csv",
    )

    # -----------------------------------------------------------------
    # 2. Participant summary used only for the accuracy plot
    # -----------------------------------------------------------------

    participant_summary = (
        merged.groupby("participant", as_index=False)
        .agg(
            Matches=("matches_answer_key", "sum"),
            Total=("matches_answer_key", "size"),
        )
    )

    participant_summary["Accuracy (%)"] = (
        100
        * participant_summary["Matches"]
        / participant_summary["Total"]
    )

    participant_summary = participant_summary.rename(
        columns={
            "participant": "Participant",
        }
    )

    participant_summary["Accuracy (%)"] = participant_summary[
        "Accuracy (%)"
    ].round(2)

    # -----------------------------------------------------------------
    # 3. Question summary
    # -----------------------------------------------------------------

    response_counts = (
        merged.groupby(QUESTION_COLUMN)["selected_option"]
        .value_counts()
        .unstack(fill_value=0)
    )

    if "A" not in response_counts.columns:
        response_counts["A"] = 0

    if "B" not in response_counts.columns:
        response_counts["B"] = 0

    response_counts = response_counts[
        ["A", "B"]
    ].reset_index()

    response_counts = response_counts.rename(
        columns={
            "A": "Selected A",
            "B": "Selected B",
        }
    )

    question_summary = (
        merged.groupby(
            [
                QUESTION_COLUMN,
                "correct_option",
            ],
            as_index=False,
        )
        .agg(
            Matches=("matches_answer_key", "sum"),
            Participants=("matches_answer_key", "size"),
        )
    )

    question_summary["Agreement with Answer Key (%)"] = (
        100
        * question_summary["Matches"]
        / question_summary["Participants"]
    )

    question_summary = question_summary.merge(
        response_counts,
        on=QUESTION_COLUMN,
        how="left",
        validate="one_to_one",
    )

    question_summary["Majority Choice"] = np.select(
        [
            question_summary["Selected A"] > question_summary["Selected B"],
            question_summary["Selected B"] > question_summary["Selected A"],
        ],
        [
            "A",
            "B",
        ],
        default="Tie",
    )

    question_summary["Question Number"] = question_summary[
        QUESTION_COLUMN
    ].apply(
        question_number
    )

    question_summary = question_summary.sort_values(
        "Question Number"
    ).reset_index(drop=True)

    question_summary_export = question_summary[
        [
            QUESTION_COLUMN,
            "correct_option",
            "Matches",
            "Participants",
            "Agreement with Answer Key (%)",
            "Selected A",
            "Selected B",
            "Majority Choice",
        ]
    ].rename(
        columns={
            QUESTION_COLUMN: "Question",
            "correct_option": "Correct Option",
            "Matches": "Matches Answer Key",
        }
    )

    question_summary_export[
        "Agreement with Answer Key (%)"
    ] = question_summary_export[
        "Agreement with Answer Key (%)"
    ].round(2)

    save_dataframe(
        question_summary_export,
        "question_summary.csv",
    )

    # -----------------------------------------------------------------
    # 4. Response matrix
    # -----------------------------------------------------------------

    response_matrix = merged.pivot(
        index=QUESTION_COLUMN,
        columns="participant",
        values="selected_option",
    )

    response_matrix = response_matrix.reindex(
        sorted(response_matrix.index, key=question_number)
    )

    response_matrix = response_matrix[
        sorted(response_matrix.columns)
    ]

    correct_options_by_question = (
        answer_key.set_index(QUESTION_COLUMN)["correct_option"]
        .reindex(response_matrix.index)
    )

    response_matrix_with_correct = response_matrix.copy()

    response_matrix_with_correct.insert(
        0,
        "Correct Option",
        correct_options_by_question,
    )

    response_matrix_export = (
        response_matrix_with_correct.reset_index()
        .rename(
            columns={
                QUESTION_COLUMN: "Question",
            }
        )
    )

    save_dataframe(
        response_matrix_export,
        "response_matrix.csv",
    )

    export_response_matrix_excel(
        response_matrix_export=response_matrix_export,
        participant_columns=list(response_matrix.columns),
    )

    # -----------------------------------------------------------------
    # 5. Pairwise participant agreement
    # -----------------------------------------------------------------

    participants = list(response_matrix.columns)

    if len(participants) != len(set(participants)):
        raise ValueError(
            "Duplicate participant IDs were found. Ensure every participant "
            "file has a unique filename, for example:\n"
            "participant_01.csv\n"
            "participant_02.csv\n"
            "participant_03.csv"
        )

    number_of_participants = len(participants)

    pairwise_values = np.full(
        (
            number_of_participants,
            number_of_participants,
        ),
        np.nan,
        dtype=float,
    )

    for row_index, participant_a in enumerate(participants):
        for column_index, participant_b in enumerate(participants):
            answers_a = response_matrix[participant_a]
            answers_b = response_matrix[participant_b]

            valid_mask = answers_a.notna() & answers_b.notna()

            if valid_mask.sum() > 0:
                pairwise_values[
                    row_index,
                    column_index,
                ] = 100 * (
                    answers_a[valid_mask].to_numpy()
                    == answers_b[valid_mask].to_numpy()
                ).mean()

    pairwise_agreement = pd.DataFrame(
        pairwise_values,
        index=participants,
        columns=participants,
    )

    pairwise_agreement.index.name = "Participant"
    pairwise_agreement.columns.name = "Participant"

    pairwise_agreement_export = pairwise_agreement.reset_index()

    save_dataframe(
        pairwise_agreement_export,
        "pairwise_agreement.csv",
    )

    # -----------------------------------------------------------------
    # 6. Plots
    # -----------------------------------------------------------------

    create_participant_accuracy_plot(
        participant_summary,
    )

    create_question_accuracy_plot(
        question_summary_export,
    )

    create_option_selection_plot(
        question_summary_export,
    )

    create_response_heatmap(
        response_matrix=response_matrix,
        correct_options_by_question=correct_options_by_question,
    )

    create_pairwise_agreement_heatmap(
        pairwise_agreement,
    )

    print("\nStudy 2 analysis completed successfully.")
    print(f"Participants found: {len(participant_ids)}")
    print(f"Questions found: {len(answer_key)}")
    print(f"Output folder: {OUTPUT_DIR.resolve()}")

    print("\nCreated files:")
    print("- all_responses_merged.csv")
    print("- question_summary.csv")
    print("- response_matrix.csv")
    print("- response_matrix.xlsx")
    print("- pairwise_agreement.csv")
    print("- participant_accuracy.png")
    print("- question_accuracy.png")
    print("- option_selection_by_question.png")
    print("- response_heatmap.png")
    print("- pairwise_agreement_heatmap.png")


if __name__ == "__main__":
    main()
